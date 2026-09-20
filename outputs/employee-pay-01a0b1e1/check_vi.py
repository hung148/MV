import openpyxl
a=openpyxl.load_workbook('C:/Users/ASUS/Downloads/MV_Manufacturing_Biweekly_Payroll.xlsx')
b=openpyxl.load_workbook('D:/mv/outputs/employee-pay-01a0b1e1/MV_Manufacturing_Bang_Luong_Tieng_Viet.xlsx')
diff=[]
for s,t in zip(a,b):
    print(s.title,'logos',len(s._images),len(t._images),'tables',len(s.tables),len(t.tables))
    for row in s:
        for c in row:
            translated=(s.title=='Employee Form' and c.column==2 and 12<=c.row<=25) or (s.title=='Payment Register' and c.column==10)
            if c.data_type=='f' and not translated and c.value!=t[c.coordinate].value:
                diff.append((s.title,c.coordinate))
print('Unexpected calculation changes:',diff)
assert not diff
